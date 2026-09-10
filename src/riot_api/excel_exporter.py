from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import CollectionError
from .parser import (
    FINAL_FIELDS,
    LEAGUE_DATA_COLUMNS,
    MASTERY_FIELDS,
    PARTICIPANT_FIELDS,
    validate_league_rows,
)
from .schemas import LeagueRowSchema

SEARCH_COLUMNS = [
    "입력한 Riot ID",
    "PUUID",
    "한국 서버 확인 결과",
    "요청 경기 수",
    "실제 수집 경기 수",
    "생성된 데이터 행 수",
    "수집 시작 시각",
    "수집 완료 시각",
    "Queue ID",
    "Platform Routing",
    "Regional Routing",
]
ERROR_COLUMNS = [
    "수집 단계",
    "PUUID",
    "Match ID",
    "Champion ID",
    "API 종류",
    "HTTP 상태 코드",
    "오류 내용",
    "재시도 횟수",
]


def unique_output_path(path: str | Path) -> Path:
    candidate = Path(path)
    if not candidate.exists():
        return candidate
    for number in range(1, 10000):
        numbered = candidate.with_name(f"{candidate.stem}_{number}{candidate.suffix}")
        if not numbered.exists():
            return numbered
    raise FileExistsError("사용 가능한 출력 파일명을 만들 수 없습니다.")


def export_workbook(
    path: str | Path,
    rows: list[dict[str, Any]],
    search_info: dict[str, Any],
    errors: list[CollectionError],
) -> Path:
    validation_errors = validate_league_rows(rows)
    if validation_errors:
        raise ValueError("엑셀 생성 전 검증 실패:\n" + "\n".join(validation_errors))
    for row in rows:
        LeagueRowSchema.model_validate(row)
    output_path = unique_output_path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    league_sheet = workbook.active
    league_sheet.title = "league_data"
    _write_table(
        league_sheet,
        LEAGUE_DATA_COLUMNS,
        [[row.get(c) for c in LEAGUE_DATA_COLUMNS] for row in rows],
    )
    _write_table(
        workbook.create_sheet("검색정보"),
        SEARCH_COLUMNS,
        [[search_info.get(c) for c in SEARCH_COLUMNS]],
    )
    _write_table(
        workbook.create_sheet("컬럼정의서"),
        ["No.", "컬럼명", "분류", "설명", "데이터 출처", "API 원본 필드"],
        _column_dictionary(),
    )
    _write_table(
        workbook.create_sheet("수집오류"),
        ERROR_COLUMNS,
        [
            [
                error.stage,
                error.puuid,
                error.match_id,
                error.champion_id,
                error.api_type,
                error.status_code,
                error.message,
                error.retries,
            ]
            for error in errors
        ],
    )
    workbook.save(output_path)

    check = load_workbook(output_path, read_only=True, data_only=False)
    try:
        if check.sheetnames[:4] != ["league_data", "검색정보", "컬럼정의서", "수집오류"]:
            raise ValueError("엑셀 생성 후 시트 검증에 실패했습니다.")
        headers = [
            cell.value for cell in next(check["league_data"].iter_rows(min_row=1, max_row=1))
        ]
        if headers != LEAGUE_DATA_COLUMNS or len(headers) != 94:
            raise ValueError("엑셀 생성 후 94개 컬럼 검증에 실패했습니다.")
        saved_rows = [
            dict(zip(headers, values, strict=True))
            for values in check["league_data"].iter_rows(min_row=2, values_only=True)
        ]
        post_errors = validate_league_rows(saved_rows)
        if post_errors:
            raise ValueError("엑셀 생성 후 데이터 검증 실패:\n" + "\n".join(post_errors))
    finally:
        check.close()
    return output_path


def _write_table(sheet: Any, headers: list[str], rows: list[list[Any]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    dark_blue = PatternFill("solid", fgColor="17365D")
    white_font = Font(color="FFFFFF", bold=True)
    gray = PatternFill("solid", fgColor="F2F2F2")
    for cell in sheet[1]:
        cell.fill = dark_blue
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_number in range(2, sheet.max_row + 1):
        if row_number % 2 == 0:
            for cell in sheet[row_number]:
                cell.fill = gray
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False
    for column_number, header in enumerate(headers, 1):
        max_length = max(
            [len(str(header))]
            + [
                len(str(sheet.cell(row, column_number).value or ""))
                for row in range(2, sheet.max_row + 1)
            ]
        )
        sheet.column_dimensions[get_column_letter(column_number)].width = min(
            max(max_length + 2, 10), 32
        )
    if "game_start_utc" in headers:
        for utc_header in ("game_start_utc", "champion_mastery_lastPlayTime_utc"):
            column = headers.index(utc_header) + 1
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column).number_format = "yyyy-mm-dd hh:mm:ss"
    if "win" in headers and sheet.max_row >= 2:
        win_column = get_column_letter(headers.index("win") + 1)
        data_range = f"A2:{get_column_letter(len(headers))}{sheet.max_row}"
        sheet.conditional_formatting.add(
            data_range,
            FormulaRule(
                formula=[f"${win_column}2=TRUE"], fill=PatternFill("solid", fgColor="DDEBF7")
            ),
        )
        sheet.conditional_formatting.add(
            data_range,
            FormulaRule(
                formula=[f"${win_column}2=FALSE"], fill=PatternFill("solid", fgColor="FCE4D6")
            ),
        )


def _column_dictionary() -> list[list[Any]]:
    descriptions: dict[str, tuple[str, str, str, str]] = {}
    meta_fields = {
        "game_id": "metadata.matchId",
        "game_start_utc": "info.gameStartTimestamp",
        "game_duration": "info.gameDuration",
        "game_mode": "info.gameMode",
        "game_type": "info.gameType",
        "game_version": "info.gameVersion",
        "map_id": "info.mapId",
        "platform_id": "info.platformId",
        "queue_id": "info.queueId",
    }
    for column, field in meta_fields.items():
        descriptions[column] = ("경기", f"경기 메타데이터 {column}", "MATCH-V5", field)
    for column, field in PARTICIPANT_FIELDS.items():
        descriptions[column] = (
            "참가자",
            f"참가자 {column}",
            "MATCH-V5",
            f"info.participants[].{field}",
        )
    descriptions["summoner_name"] = (
        "참가자",
        "Riot ID 게임 이름, 없으면 소환사 이름",
        "MATCH-V5",
        "info.participants[].riotIdGameName / summonerName",
    )
    for column in ("solo_tier", "solo_rank", "solo_lp", "solo_wins", "solo_losses"):
        field = {"solo_lp": "leaguePoints"}.get(column, column.removeprefix("solo_"))
        descriptions[column] = ("솔로랭크", "현재 솔로랭크 정보", "LEAGUE-V4", field)
    for column in ("flex_tier", "flex_rank", "flex_lp", "flex_wins", "flex_losses"):
        descriptions[column] = ("호환", "스키마 호환용 빈 컬럼", "미수집", "")
    for column, field in MASTERY_FIELDS.items():
        descriptions[column] = ("숙련도", "챔피언 숙련도 정보", "CHAMPION-MASTERY-V4", field)
    descriptions["champion_mastery_lastPlayTime_utc"] = (
        "숙련도",
        "lastPlayTime의 UTC 변환값",
        "CHAMPION-MASTERY-V4",
        "lastPlayTime",
    )
    for column, field in FINAL_FIELDS.items():
        descriptions[column] = (
            "최종스탯",
            "타임라인 마지막 유효 프레임의 최종 스탯",
            "MATCH-V5 Timeline",
            f"championStats.{field}",
        )
    return [
        [number, column, *(descriptions.get(column) or ("참가자", column, "MATCH-V5", column))]
        for number, column in enumerate(LEAGUE_DATA_COLUMNS, 1)
    ]
