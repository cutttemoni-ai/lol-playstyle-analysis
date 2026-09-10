from __future__ import annotations

from datetime import datetime

from openpyxl import load_workbook

from riot_api.excel_exporter import ERROR_COLUMNS, SEARCH_COLUMNS, export_workbook
from riot_api.models import CollectionError
from riot_api.parser import LEAGUE_DATA_COLUMNS


def valid_rows():
    rows = []
    for participant_id in range(1, 11):
        row = {column: None for column in LEAGUE_DATA_COLUMNS}
        row.update(
            {
                "game_id": "KR_1",
                "game_start_utc": datetime(2026, 1, 1),
                "platform_id": "KR",
                "queue_id": 420,
                "participant_id": participant_id,
                "puuid": f"p-{participant_id}",
                "win": participant_id <= 5,
            }
        )
        rows.append(row)
    return rows


def test_workbook_has_required_sheets_and_exact_schema(tmp_path):
    output = export_workbook(tmp_path / "result.xlsx", valid_rows(), {}, [])
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["league_data", "검색정보", "컬럼정의서", "수집오류"]
    assert [cell.value for cell in workbook["league_data"][1]] == LEAGUE_DATA_COLUMNS
    assert workbook["league_data"].max_column == 94
    assert workbook["컬럼정의서"].max_row == 95
    assert [cell.value for cell in workbook["검색정보"][1]] == SEARCH_COLUMNS
    assert [cell.value for cell in workbook["수집오류"][1]] == ERROR_COLUMNS


def test_partial_api_failure_still_creates_workbook(tmp_path):
    error = CollectionError(stage="타임라인", match_id="KR_1", api_type="MATCH-V5", status_code=503)
    output = export_workbook(tmp_path / "partial.xlsx", valid_rows(), {}, [error])
    workbook = load_workbook(output)
    assert workbook["league_data"].max_row == 11
    assert workbook["수집오류"].max_row == 2
    assert all(workbook["league_data"].cell(row, 70).value is None for row in range(2, 12))


def test_existing_output_is_not_overwritten(tmp_path):
    original = tmp_path / "result.xlsx"
    original.write_bytes(b"keep")
    output = export_workbook(original, valid_rows(), {}, [])
    assert output.name == "result_1.xlsx"
    assert original.read_bytes() == b"keep"
